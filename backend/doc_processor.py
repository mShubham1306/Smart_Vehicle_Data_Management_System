import os
import re
import uuid
import shutil
import logging
from io import BytesIO
from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional

from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Query, BackgroundTasks
from fastapi.responses import StreamingResponse, FileResponse
from bson import ObjectId
import pdfplumber
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

try:
    from PIL import Image
    import pytesseract
except ImportError:
    Image = None
    pytesseract = None

from database import doc_uploads_collection, users_collection
from auth import get_current_user, require_admin

# Setup router
router = APIRouter()

# Setup logging
logger = logging.getLogger("doc_processor")
logger.setLevel(logging.INFO)

# Directory for storing document uploads
UPLOAD_DIR = "uploads_docs"
os.makedirs(UPLOAD_DIR, exist_ok=True)

# Helper to serialize MongoDB documents
def serialize_doc(doc: Dict[str, Any]) -> Dict[str, Any]:
    serialized = dict(doc)
    if "_id" in serialized:
        serialized["id"] = str(serialized["_id"])
        del serialized["_id"]
    if "upload_time" in serialized and isinstance(serialized["upload_time"], datetime):
        serialized["upload_time"] = serialized["upload_time"].isoformat()
    return serialized

# Helper to extract fields using regular expressions
def extract_fields_from_text(text: str) -> Dict[str, str]:
    fields = {
        "insured_name": "",
        "policy_number": "",
        "vehicle_number": "",
        "vehicle_model": "",
        "engine_number": "",
        "chassis_number": "",
        "insurance_company": "Unknown / Detected",
        "policy_start_date": "",
        "policy_end_date": "",
        "idv": "",
        "premium": "",
        "mobile_number": "",
        "email": "",
        "rto": ""
    }

    if not text:
        return fields

    # 1. Insurance Company Detection
    text_lower = text.lower()
    if "icici" in text_lower or "lombard" in text_lower:
        fields["insurance_company"] = "ICICI Lombard"
    elif "hdfc" in text_lower or "ergo" in text_lower:
        fields["insurance_company"] = "HDFC ERGO"
    elif "bajaj" in text_lower or "allianz" in text_lower:
        fields["insurance_company"] = "Bajaj Allianz"
    elif "tata" in text_lower or "aig" in text_lower:
        fields["insurance_company"] = "TATA AIG"
    elif "reliance" in text_lower:
        fields["insurance_company"] = "Reliance General"
    elif "iffco" in text_lower or "tokio" in text_lower:
        fields["insurance_company"] = "IFFCO Tokio"
    elif "national" in text_lower and "insurance" in text_lower:
        fields["insurance_company"] = "National Insurance"
    elif "new india" in text_lower or "assurance" in text_lower:
        fields["insurance_company"] = "New India Assurance"
    elif "oriental" in text_lower:
        fields["insurance_company"] = "Oriental Insurance"
    elif "united india" in text_lower:
        fields["insurance_company"] = "United India Insurance"
    elif "chola" in text_lower or "ms" in text_lower:
        fields["insurance_company"] = "Cholamandalam MS"

    # 2. Vehicle Number (Indian registration plates)
    # Formats: MH02CL1234, GJ-15-EK-8917, DL 3C AY 1122, etc.
    veh_match = re.search(r'\b([A-Z]{2}[-\s]?\d{2}[-\s]?[A-Z]{1,3}[-\s]?\d{4})\b', text.upper())
    if veh_match:
        # Clean delimiters
        fields["vehicle_number"] = re.sub(r'[-\s]', '', veh_match.group(1))

    # 3. Policy Number
    # Try generic keywords or specific formats
    policy_patterns = [
        r'(?i)(?:policy\s*(?:no|number)|policy\s*id)\s*[:\-\s\.]*\s*([A-Z0-9/\\\-_]+)',
        r'\b(3005[/\\]\d+[/\\]\d+[/\\]\d+[/\\]\d+)\b', # ICICI Lombard typical pattern
        r'\b([A-Z0-9]{2,}[/\\][A-Z0-9/\\\-_]{8,})\b'
    ]
    for pattern in policy_patterns:
        match = re.search(pattern, text)
        if match:
            # Clean matching group and make sure it has digits (not just words)
            val = match.group(1).strip()
            if any(c.isdigit() for c in val) and len(val) > 5:
                fields["policy_number"] = val
                break

    # 4. Dates (Start and End Dates)
    # Search for all dates formatted as DD/MM/YYYY or DD-MM-YYYY or YYYY-MM-DD
    dates = re.findall(r'\b(\d{2}[-/\.]\d{2}[-/\.]\d{4})\b', text)
    if not dates:
        # Check YYYY-MM-DD
        dates_alt = re.findall(r'\b(\d{4}[-/\.]\d{2}[-/\.]\d{2})\b', text)
        if dates_alt:
            # Convert to DD-MM-YYYY format
            dates = []
            for d in dates_alt:
                parts = re.split(r'[-/\.]', d)
                dates.append(f"{parts[2]}-{parts[1]}-{parts[0]}")

    if dates:
        # Often start date is the first found, and end date is the second found,
        # but let's be smarter if we see keywords
        fields["policy_start_date"] = dates[0]
        if len(dates) > 1:
            fields["policy_end_date"] = dates[1]

    # Date keyword search override
    start_match = re.search(r'(?i)(?:period\s*of\s*insurance\s*from|duration\s*from|valid\s*from|effective\s*date|start\s*date)[^\d]*(\d{2}[-/\.]\d{2}[-/\.]\d{4})', text)
    if start_match:
        fields["policy_start_date"] = start_match.group(1)

    end_match = re.search(r'(?i)(?:period\s*of\s*insurance\s*to|duration\s*to|valid\s*to|expiry\s*date|end\s*date|expires\s*on)[^\d]*(\d{2}[-/\.]\d{2}[-/\.]\d{4})', text)
    if end_match:
        fields["policy_end_date"] = end_match.group(1)

    # 5. Customer Name (Insured Name)
    # Often comes after "Insured Name", "Name of Insured", "Customer Name", "Mr/Mrs/Ms"
    name_match = re.search(r'(?i)(?:insured[\s\']*(?:name|details)|name\s*of\s*insured|customer\s*name)\s*[:\-\s\.]*\s*(?:mr|mrs|ms|m/s)?[\.\s]*([A-Z][A-Z\s]{2,30})', text)
    if name_match:
        fields["insured_name"] = name_match.group(1).strip()
    else:
        # Fallback search for names in capitalization
        name_fallback = re.search(r'(?i)insured\s*:\s*([A-Z][A-Z\s]{2,25})', text)
        if name_fallback:
            fields["insured_name"] = name_fallback.group(1).strip()

    # 6. IDV
    idv_match = re.search(r'(?i)(?:idv|insured\s*declared\s*value)[^\d₹]*([\d,]{3,9})', text)
    if idv_match:
        fields["idv"] = idv_match.group(1).replace(",", "").strip()

    # 7. Premium
    prem_match = re.search(r'(?i)(?:net\s*premium|total\s*premium|premium\s*payable|final\s*premium)[^\d₹]*([\d,]{3,9})', text)
    if prem_match:
        fields["premium"] = prem_match.group(1).replace(",", "").strip()
    else:
        # Try generic premium search
        prem_generic = re.search(r'(?i)premium[^\d₹]*([\d,]{3,9})', text)
        if prem_generic:
            fields["premium"] = prem_generic.group(1).replace(",", "").strip()

    # 8. Engine Number & Chassis Number
    engine_match = re.search(r'(?i)(?:engine\s*(?:no|number))\s*[:\-\s\.]*\s*([A-Z0-9]+)\b', text)
    if engine_match:
        fields["engine_number"] = engine_match.group(1).strip()

    chassis_match = re.search(r'(?i)(?:chassis\s*(?:no|number))\s*[:\-\s\.]*\s*([A-Z0-9]{5,20})\b', text)
    if chassis_match:
        fields["chassis_number"] = chassis_match.group(1).strip()

    # 9. Vehicle Model
    model_match = re.search(r'(?i)(?:vehicle\s*model|model|make\s*&\s*model|vehicle\s*make)[:\-\s\.]*([A-Z0-9][A-Z0-9\s]{2,30})', text)
    if model_match:
        fields["vehicle_model"] = model_match.group(1).strip()

    # 10. Mobile Number
    mobile_match = re.search(r'\b([6-9]\d{9})\b', text)
    if mobile_match:
        fields["mobile_number"] = mobile_match.group(1)

    # 11. Email
    email_match = re.search(r'\b([a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})\b', text)
    if email_match:
        fields["email"] = email_match.group(1)

    # 12. RTO (State code + RTO code, e.g. GJ-15 or GJ15)
    rto_match = re.search(r'\b([A-Z]{2}[-\s]?\d{2})\b', text.upper())
    if rto_match:
        fields["rto"] = re.sub(r'[-\s]', '', rto_match.group(1))

    return fields


def run_ocr_on_file(file_path: str, file_type: str) -> tuple[str, float]:
    """
    Runs text extraction.
    Returns (extracted_text, confidence_score).
    """
    text = ""
    confidence = 0.8  # Default confidence

    if file_type == "pdf":
        try:
            with pdfplumber.open(file_path) as pdf:
                pages_text = []
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        pages_text.append(page_text)
                text = "\n".join(pages_text)
                if text.strip():
                    return text, 0.95  # Digital PDF gets high confidence
        except Exception as e:
            logger.error(f"pdfplumber failed: {e}")

    # Fallback to image OCR if pdfplumber failed (scanned PDF) or if it is an image
    if Image and pytesseract:
        try:
            # If it is a PDF that failed or had no text, we would need pdf2image.
            # As fallback, we only support direct image files.
            if file_type in ["png", "jpg", "jpeg"]:
                img = Image.open(file_path)
                text = pytesseract.image_to_string(img)
                confidence = 0.85
        except Exception as e:
            logger.error(f"pytesseract failed: {e}")

    return text, confidence


# --- ROUTES ---

@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    Upload a document, perform OCR, parse fields, and save to MongoDB.
    """
    # Verify file extension
    ext = file.filename.split(".")[-1].lower() if "." in file.filename else ""
    if ext not in ["pdf", "png", "jpg", "jpeg"]:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file type. Only PDF, PNG, JPG, and JPEG are allowed."
        )

    # Generate unique document ID
    doc_id = str(uuid.uuid4())
    file_name = f"{doc_id}_{file.filename}"
    file_path = os.path.join(UPLOAD_DIR, file_name)

    # Save file to disk
    start_time = datetime.utcnow()
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        logger.error(f"Failed to write file to disk: {e}")
        raise HTTPException(status_code=500, detail="Could not save file to disk.")

    # Perform OCR
    extracted_text, ocr_confidence = run_ocr_on_file(file_path, ext)
    processing_time_ms = int((datetime.utcnow() - start_time).total_seconds() * 1000)

    # Parse fields
    extracted_fields = extract_fields_from_text(extracted_text)

    # Check status
    if not extracted_text.strip():
        ocr_status = "failed"
    elif not any(extracted_fields.values()):
        ocr_status = "partial"
    else:
        ocr_status = "success"

    # Get user name and email from database
    user_id = current_user["id"]
    db_user = await users_collection.find_one({"_id": ObjectId(user_id)})
    user_name = db_user.get("name", db_user.get("username", "User")) if db_user else current_user.get("username", "User")
    user_email = db_user.get("email", "") if db_user else ""

    # Document schema
    doc_data = {
        "doc_id": doc_id,
        "uploaded_by_user_id": user_id,
        "uploaded_by_name": user_name,
        "uploaded_by_email": user_email,
        "upload_time": datetime.utcnow(),
        "file_name": file.filename,
        "file_type": ext,
        "file_path": file_path,
        "ocr_status": ocr_status,
        "ocr_confidence": ocr_confidence,
        "processing_time_ms": processing_time_ms,
        "extracted_fields": extracted_fields,
        "raw_text_snippet": extracted_text[:1000] # Save a snippet for debugging
    }

    # Insert to MongoDB
    await doc_uploads_collection.insert_one(doc_data)

    return serialize_doc(doc_data)


@router.get("/")
async def list_documents(
    search: Optional[str] = Query(None, description="Search term for vehicle, policy, customer"),
    ocr_status: Optional[str] = Query(None, description="Filter by OCR status"),
    user_id: Optional[str] = Query(None, description="Filter by user (Admin only)"),
    page: int = Query(1, ge=1),
    limit: int = Query(10, ge=1, le=100),
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    List documents uploaded. Users see only their own uploads, admins see all.
    """
    query: Dict[str, Any] = {}

    # Enforce role scoping
    if current_user.get("role") != "admin":
        query["uploaded_by_user_id"] = current_user["id"]
    elif user_id and user_id != "all":
        query["uploaded_by_user_id"] = user_id

    # Apply filters
    if ocr_status:
        query["ocr_status"] = ocr_status

    if search:
        search_reg = re.compile(search, re.IGNORECASE)
        query["$or"] = [
            {"file_name": search_reg},
            {"extracted_fields.vehicle_number": search_reg},
            {"extracted_fields.policy_number": search_reg},
            {"extracted_fields.insured_name": search_reg},
            {"extracted_fields.insurance_company": search_reg}
        ]

    # Query DB
    skip = (page - 1) * limit
    cursor = doc_uploads_collection.find(query).sort("upload_time", -1).skip(skip).limit(limit)
    total_docs = await doc_uploads_collection.count_documents(query)

    docs = await cursor.to_list(length=limit)
    serialized_docs = [serialize_doc(d) for d in docs]

    # Fetch distinct users who have uploaded docs (for admin filter dropdown)
    users_list = []
    if current_user.get("role") == "admin":
        # Get unique user IDs from doc uploads
        pipeline = [
            {"$group": {"_id": "$uploaded_by_user_id", "name": {"$first": "$uploaded_by_name"}}}
        ]
        user_cursor = doc_uploads_collection.aggregate(pipeline)
        users_list = [{"id": u["_id"], "name": u["name"]} async for u in user_cursor if u["_id"]]

    return {
        "documents": serialized_docs,
        "total": total_docs,
        "page": page,
        "limit": limit,
        "users": users_list
    }


@router.get("/{doc_id}")
async def get_document(
    doc_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    Get detailed document information.
    """
    query = {"doc_id": doc_id}
    if current_user.get("role") != "admin":
        query["uploaded_by_user_id"] = current_user["id"]

    doc = await doc_uploads_collection.find_one(query)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    return serialize_doc(doc)


@router.get("/{doc_id}/download")
async def download_document_file(
    doc_id: str,
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    Download the original uploaded document file.
    """
    query = {"doc_id": doc_id}
    if current_user.get("role") != "admin":
        query["uploaded_by_user_id"] = current_user["id"]

    doc = await doc_uploads_collection.find_one(query)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    file_path = doc.get("file_path")
    if not file_path or not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Physical file not found on disk.")

    return FileResponse(file_path, filename=doc.get("file_name"))


@router.patch("/{doc_id}")
async def update_document_fields(
    doc_id: str,
    payload: Dict[str, Any],
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    Manually correct or update the extracted fields of a document.
    """
    query = {"doc_id": doc_id}
    if current_user.get("role") != "admin":
        query["uploaded_by_user_id"] = current_user["id"]

    doc = await doc_uploads_collection.find_one(query)
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    # Only allow updating fields inside extracted_fields
    updates = {}
    if "extracted_fields" in payload:
        for k, v in payload["extracted_fields"].items():
            updates[f"extracted_fields.{k}"] = str(v).strip()

    if "ocr_status" in payload:
        updates["ocr_status"] = payload["ocr_status"]

    if not updates:
        raise HTTPException(status_code=400, detail="No valid fields provided for update.")

    await doc_uploads_collection.update_one(query, {"$set": updates})

    # Fetch updated document
    updated_doc = await doc_uploads_collection.find_one(query)
    return serialize_doc(updated_doc)


@router.delete("/{doc_id}")
async def delete_document(
    doc_id: str,
    current_user: Dict[str, Any] = Depends(require_admin)
):
    """
    Delete a document record and its corresponding file on disk (Admin only).
    """
    doc = await doc_uploads_collection.find_one({"doc_id": doc_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")

    # Remove local file if it exists
    file_path = doc.get("file_path")
    if file_path and os.path.exists(file_path):
        try:
            os.remove(file_path)
        except Exception as e:
            logger.error(f"Failed to delete file {file_path}: {e}")

    await doc_uploads_collection.delete_one({"doc_id": doc_id})
    return {"message": "Document deleted successfully."}


@router.get("/export/excel")
async def export_documents_excel(
    user_id: Optional[str] = Query("all", description="Export specific user or 'all' for admin"),
    current_user: Dict[str, Any] = Depends(require_admin)
):
    """
    Export processed documents to a formatted Excel file (Admin only).
    Supports filtering by specific user or exporting all.
    """
    query: Dict[str, Any] = {}
    if user_id and user_id != "all":
        query["uploaded_by_user_id"] = user_id

    # Fetch all records matching query
    cursor = doc_uploads_collection.find(query).sort("upload_time", -1)
    docs = await cursor.to_list(length=10000)

    # Create openpyxl Workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Processed Documents"

    # Setup Styles
    title_font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Calibri", size=11, bold=False)
    
    title_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    even_row_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Sheet Title block
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"SmartInsure Enterprise Document Processing Report — {datetime.now().strftime('%Y-%m-%d')}"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # 2. Table Headers
    headers = [
        "Uploaded By", "User ID", "Upload Date", "File Name", 
        "Policy No", "Vehicle No", "Customer Name", "Premium", 
        "IDV", "Start Date", "End Date", "OCR Status"
    ]
    
    ws.row_dimensions[3].height = 25
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 3. Fill Data Rows
    row_idx = 4
    for doc in docs:
        extracted = doc.get("extracted_fields", {})
        
        # Format upload date
        up_time = doc.get("upload_time")
        up_time_str = up_time.strftime("%Y-%m-%d %H:%M:%S") if isinstance(up_time, datetime) else ""
        
        # Row data array
        row_data = [
            doc.get("uploaded_by_name", ""),
            doc.get("uploaded_by_user_id", ""),
            up_time_str,
            doc.get("file_name", ""),
            extracted.get("policy_number", ""),
            extracted.get("vehicle_number", ""),
            extracted.get("insured_name", ""),
            extracted.get("premium", ""),
            extracted.get("idv", ""),
            extracted.get("policy_start_date", ""),
            extracted.get("policy_end_date", ""),
            doc.get("ocr_status", "unknown").upper()
        ]

        ws.row_dimensions[row_idx].height = 20
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.font = body_font
            cell.border = thin_border
            
            # Alternate row coloring
            if row_idx % 2 == 0:
                cell.fill = even_row_fill
                
            # Alignments
            if col_idx in [3, 8, 9, 10, 11, 12]:  # Dates, numbers, status
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        row_idx += 1

    # 4. Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        # Skip checking the merged row 1 length
        for cell in col[2:]: 
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save output to a stream
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"SmartInsure_Documents_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/analytics")
async def get_analytics_dashboard(
    current_user: Dict[str, Any] = Depends(require_admin)
):
    """
    Get document processing analytics (Admin only).
    Returns total count, today's uploads, month's uploads, OCR stats, and user counts.
    """
    # 1. Date ranges
    now = datetime.utcnow()
    today_start = datetime(now.year, now.month, now.day)
    month_start = datetime(now.year, now.month, 1)

    # 2. Aggregations
    total_uploads = await doc_uploads_collection.count_documents({})
    today_uploads = await doc_uploads_collection.count_documents({"upload_time": {"$gte": today_start}})
    month_uploads = await doc_uploads_collection.count_documents({"upload_time": {"$gte": month_start}})

    # Status counts
    success_count = await doc_uploads_collection.count_documents({"ocr_status": "success"})
    partial_count = await doc_uploads_collection.count_documents({"ocr_status": "partial"})
    failed_count = await doc_uploads_collection.count_documents({"ocr_status": "failed"})

    success_rate = (success_count / total_uploads * 100) if total_uploads > 0 else 0.0

    # User-wise counts (Productivity tracking)
    pipeline = [
        {
            "$group": {
                "_id": "$uploaded_by_user_id",
                "username": {"$first": "$uploaded_by_name"},
                "email": {"$first": "$uploaded_by_email"},
                "total": {"$sum": 1},
                "success": {"$sum": {"$cond": [{"$eq": ["$ocr_status", "success"]}, 1, 0]}},
                "failed": {"$sum": {"$cond": [{"$eq": ["$ocr_status", "failed"]}, 1, 0]}},
                "last_active": {"$max": "$upload_time"}
            }
        },
        {"$sort": {"total": -1}}
    ]
    user_cursor = doc_uploads_collection.aggregate(pipeline)
    user_stats = []
    async for u in user_cursor:
        if not u["_id"]:
            continue
        user_stats.append({
            "user_id": u["_id"],
            "name": u["username"] or "Unknown",
            "email": u["email"] or "",
            "total_processed": u["total"],
            "success_count": u["success"],
            "failed_count": u["failed"],
            "success_rate": (u["success"] / u["total"] * 100) if u["total"] > 0 else 0.0,
            "last_active": u["last_active"].isoformat() if isinstance(u["last_active"], datetime) else ""
        })

    return {
        "total_uploads": total_uploads,
        "today_uploads": today_uploads,
        "month_uploads": month_uploads,
        "success_rate": round(success_rate, 2),
        "status_distribution": {
            "success": success_count,
            "partial": partial_count,
            "failed": failed_count
        },
        "user_productivity": user_stats
    }


@router.get("/my/stats")
async def get_my_upload_stats(
    current_user: Dict[str, Any] = Depends(get_current_user)
):
    """
    Get uploading stats for the currently logged in user.
    """
    user_id = current_user["id"]
    total = await doc_uploads_collection.count_documents({"uploaded_by_user_id": user_id})
    success = await doc_uploads_collection.count_documents({"uploaded_by_user_id": user_id, "ocr_status": "success"})
    partial = await doc_uploads_collection.count_documents({"uploaded_by_user_id": user_id, "ocr_status": "partial"})
    failed = await doc_uploads_collection.count_documents({"uploaded_by_user_id": user_id, "ocr_status": "failed"})

    return {
        "total": total,
        "success": success,
        "partial": partial,
        "failed": failed,
        "success_rate": round((success / total * 100), 2) if total > 0 else 0.0
    }
